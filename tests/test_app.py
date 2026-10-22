import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

import analysis as an
import app as webapp
import ncbi
import tmpred

FIXTURE = (Path(__file__).parent / "fixtures" / "Q26473.gp").read_text()

# Residues 640-660 of Semaphorin-7A (O75326), around its motif.
SEMA7A_SNIPPET = "LLGHACALAASLWLGVLPTLT"

# Values read off the example KD plot in the problem statement (x100),
# for the 13 windows around ALHAGLW 297-303.
EXAMPLE_PLOT = [36, 49, 72, 45, 61, 51, 92, 121, 112, 149, 174, 167, 181]


@pytest.fixture
def client(monkeypatch):
    def fake_fetch(accession):
        accession = ncbi.clean_accession(accession).upper()
        if accession != "Q26473":
            raise ValueError(f"NCBI has no protein record for '{accession}'.")
        sequence = ncbi.parse_origin(FIXTURE)
        return ncbi.ProteinRecord("Q26473", FIXTURE, sequence, "Semaphorin-1A", "Schistocerca americana", "Q26473.1")

    monkeypatch.setattr(ncbi, "fetch_record", fake_fetch)
    webapp.app.config["TESTING"] = True
    return webapp.app.test_client()


# analysis

def test_example_protein_matches_assignment_plot():
    result = an.analyze(webapp.EXAMPLE_SEQUENCE)
    hit = next(h for h in result.hits if h.motif == "ALHAGLW")
    assert (hit.start, hit.end) == (297, 303)
    assert hit.qualifying
    assert len(hit.windows) == 13
    scores = [round(w.score * 100) for w in hit.windows]
    assert all(abs(a - b) <= 2 for a, b in zip(scores, EXAMPLE_PLOT))


def test_window_centre_and_running_mean():
    windows = an.sliding_windows("AAAAAIIIII", 5, 1.6)
    assert [w.center for w in windows] == [3, 4, 5, 6, 7, 8]
    assert windows[0].score == pytest.approx(1.8)
    assert windows[-1].score == pytest.approx(4.5)


def test_motif_needs_a_whole_high_window():
    # surrounded by I, so the windows around the motif are hydrophobic
    seq = "I" * 12 + "ALAAALW" + "I" * 12
    result = an.analyze(seq, an.Settings(window=19))
    assert [h.qualifying for h in result.hits] == [True]

    hydrophilic = "K" * 12 + "ALAAALW" + "K" * 12
    result = an.analyze(hydrophilic, an.Settings(window=19))
    assert result.hits and not result.hits[0].qualifying


def test_threshold_is_inclusive():
    seq = "A" * 19  # mean exactly 1.8
    assert an.analyze(seq, an.Settings(threshold=1.8)).high_window_count == 1
    assert an.analyze(seq, an.Settings(threshold=1.81)).high_window_count == 0


def test_overlapping_motifs_and_allowed_set():
    regex, length, _ = an.compile_pattern("AL???LW")
    assert length == 7
    assert [m[1] for m in an.find_motifs("ALAAALWALAALWLW", regex)] == [1]
    # E, D, K and R aren't in the allowed list
    assert an.find_motifs("ALEAALW", regex) == []
    assert an.find_motifs("ALGHQLW", regex)
    overlapping, _, _ = an.compile_pattern("A?A")
    assert [m[1] for m in an.find_motifs("AAAAA", overlapping)] == [1, 2, 3]


def test_custom_pattern_syntax():
    regex, length, readable = an.compile_pattern("n[ST]?", "AG")
    assert length == 3 and readable == "N[ST][AG]"
    assert [m[0] for m in an.find_motifs("NSANTGNSK", regex)] == ["NSA", "NTG"]
    with pytest.raises(ValueError):
        an.compile_pattern("AL[??]LW")
    with pytest.raises(ValueError):
        an.compile_pattern("AL1LW")


def test_normalize_sequence_formats():
    fasta = ">sp|X| test\nMKT LLV\n  10 aiil*\n"
    assert an.normalize_sequence(fasta)[0] == "MKTLLVAIIL"
    seq, warnings = an.normalize_sequence("MKXB")
    assert seq == "MKXB" and len(warnings) == 2
    with pytest.raises(ValueError):
        an.normalize_sequence("MKT-LLV")
    with pytest.raises(ValueError):
        an.normalize_sequence("   ")
    with pytest.raises(ValueError):
        an.normalize_sequence(">a\nMK\n>b\nMK")


def test_short_sequence_and_window_limits():
    result = an.analyze("ALAAALW")
    assert result.windows == [] and result.hits and not result.hits[0].windows
    with pytest.raises(ValueError):
        an.analyze("A" * 50, an.Settings(window=5))  # motif longer than window
    with pytest.raises(ValueError):
        an.analyze("A" * 50, an.Settings(threshold=9))


def test_segments_merge_overlapping_windows():
    seq = "K" * 20 + "I" * 25 + "K" * 20
    segments = an.analyze(seq).segments
    assert len(segments) == 1
    assert segments[0].start <= 21 and segments[0].end >= 45


# TMpred-style prediction

def test_tmpred_curves_match_the_example_plot():
    seq = ncbi.parse_origin(FIXTURE)
    io, oi = tmpred.profile(seq)
    target = json.loads((Path(__file__).parent / "fixtures" / "tmpred_example_Q26473.json").read_text())["mid"]
    rows = range(9, 715)
    ours = [(io[i] + oi[i]) / 2 for i in rows]
    theirs = [target[i] for i in rows]
    mean_a, mean_b = sum(ours) / len(ours), sum(theirs) / len(theirs)
    cov = sum((a - mean_a) * (b - mean_b) for a, b in zip(ours, theirs))
    r = cov / (sum((a - mean_a) ** 2 for a in ours) * sum((b - mean_b) ** 2 for b in theirs)) ** 0.5
    assert r > 0.95


def test_tmpred_helices_for_q26473():
    io, oi = tmpred.profile(ncbi.parse_origin(FIXTURE))
    found = tmpred.helices(io, oi)
    centers = [h["center"] for h in found]
    # signal peptide at the start (i->o) and the membrane helix near 640 (o->i), like TMpred
    assert any(c < 25 for c in centers) and any(630 <= c <= 650 for c in centers)
    assert found[0]["orientation"] == "i→o" and found[-1]["orientation"] == "o→i"


def test_tmpred_is_zero_where_the_core_does_not_fit():
    io, oi = tmpred.profile("ALAAALW" * 5)
    assert io[0] == oi[0] == 0 and io[-1] == oi[-1] == 0 and any(io)
    assert tmpred.profile("MKTA") == ([0, 0, 0, 0], [0, 0, 0, 0])


# NCBI parsing

def test_parse_origin_and_fields():
    assert len(ncbi.parse_origin(FIXTURE)) == 730
    fields = dict(ncbi.top_level_fields(FIXTURE))
    assert fields["ACCESSION"] == "Q26473"
    assert fields["LOCUS"].startswith("SEM1A_SCHAM")


def test_parse_features_handles_wrapped_qualifiers():
    features = ncbi.parse_features(FIXTURE)
    domain = next(f for f in features if f.location == "28..490")
    note = dict(domain.qualifiers)["note"]
    assert note == "Sema. /evidence=ECO:0000255|PROSITE-ProRule:PRU00352."


def test_readable_definition():
    assert ncbi.readable_definition("RecName: Full=Semaphorin-1A; AltName: Full=Fasciclin IV; Flags: Precursor.") == "Semaphorin-1A"
    assert ncbi.readable_definition("hemoglobin subunit beta [Homo sapiens].") == "hemoglobin subunit beta [Homo sapiens]"


def test_clean_accession_rejects_junk():
    with pytest.raises(ValueError):
        ncbi.clean_accession("Q26473; DROP")
    assert ncbi.clean_accession(" NP_000509.1 ") == "NP_000509.1"


# routes

def test_index(client):
    page = client.get("/").get_data(as_text=True)
    assert "Analyze" in page and "AL???LW" in page and 'name="query"' in page


def test_analyze_pasted_example(client):
    response = client.post("/analyze", data={"mode": "sequence", "sequence": webapp.EXAMPLE_SEQUENCE})
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "ALHAGLW" in page and "inside a hydrophobic window" in page
    assert 'class="motif' in page and 'class="high' in page


def test_single_box_detects_accession_or_sequence(client):
    page = client.post("/analyze", data={"query": " Q26473 "}).get_data(as_text=True)
    assert "/export/Q26473" in page
    page = client.post("/analyze", data={"query": webapp.EXAMPLE_SEQUENCE[:200]}).get_data(as_text=True)
    assert "Pasted sequence" in page
    assert webapp.looks_like_accession("NP_000509.1")
    assert not webapp.looks_like_accession("MKTLLVAIIL")


def test_fetch_returns_results_fragment(client):
    headers = {"X-Requested-With": "fetch"}
    response = client.post("/analyze", data={"query": "Q26473"}, headers=headers)
    fragment = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "<html" not in fragment and "Hydropathy profile" in fragment and "analysis-data" in fragment
    response = client.post("/analyze", data={"query": "ZZ999999"}, headers=headers)
    assert response.status_code == 400 and 'class="error"' in response.get_data(as_text=True)


def test_full_page_keeps_form_and_results_together(client):
    # no-JS fallback: normal form post returns the whole page
    page = client.post("/analyze", data={"query": "Q26473"}).get_data(as_text=True)
    assert 'id="analyze-form"' in page and "Hydropathy profile" in page


def test_results_only_after_analyze_is_pressed(client):
    page = client.get("/").get_data(as_text=True)
    assert 'class="results"' not in page and "is-idle" in page
    response = client.get("/analyze?query=Q26473")
    assert response.status_code == 302 and response.headers["Location"].endswith("/")


def test_example_chip_motif_matches():
    assert [h[0] for h in an.find_motifs(SEMA7A_SNIPPET, an.compile_pattern("AL???LW")[0])] == ["ALAASLW"]
    assert [p[1] for p in webapp.EXAMPLE_PROTEINS] == ["O75326", "P69681"]


def test_every_missing_section_explains_why(client):
    fetch = {"X-Requested-With": "fetch"}
    # No motif at all (Q26473).
    page = client.post("/analyze", data={"query": "Q26473"}, headers=fetch).get_data(as_text=True)
    assert "does not occur anywhere" in page and "Nothing to plot" in page
    assert "No letters are blue because the motif does not occur" in page
    # Motif present but below the threshold.
    page = client.post("/analyze", data={"query": webapp.EXAMPLE_SEQUENCE, "threshold": "2.1"}, headers=fetch).get_data(as_text=True)
    assert "is present, but not in a hydrophobic region" in page and "below the 2.1 threshold" in page
    # Too short to score.
    page = client.post("/analyze", data={"query": "MKALAAALWK"}, headers=fetch).get_data(as_text=True)
    assert "cannot be scored" in page and "no profile can be drawn" in page and "no scores to plot" in page
    # Nothing above the threshold anywhere.
    page = client.post("/analyze", data={"query": "K" * 40}, headers=fetch).get_data(as_text=True)
    assert "no region is marked as hydrophobic" in page and "No letters are red" in page
    # pasted sequence has no NCBI link, page should say so
    assert "No NCBI record, so the export contains the analysis only" in page


def test_results_include_tm_prediction(client):
    page = client.post("/analyze", data={"query": "Q26473"}, headers={"X-Requested-With": "fetch"}).get_data(as_text=True)
    assert "Transmembrane prediction" in page and 'id="tm-chart"' in page and "o→i" in page
    short = client.post("/analyze", data={"query": "MKALAAALWK"}, headers={"X-Requested-With": "fetch"}).get_data(as_text=True)
    assert "too short for the 21-residue helix core" in short


def test_analyze_accession_links(client):
    response = client.post("/analyze", data={"mode": "accession", "accession": "q26473"})
    page = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "https://www.ncbi.nlm.nih.gov/protein/?term=Q26473" in page
    assert "/export/Q26473" in page


def test_analyze_errors_are_friendly(client):
    response = client.post("/analyze", data={"mode": "accession", "accession": "ZZ999999"})
    assert response.status_code == 400 and "no protein record" in response.get_data(as_text=True)
    response = client.post("/analyze", data={"mode": "sequence", "sequence": "MKT", "window": "abc"})
    assert response.status_code == 400 and "whole number" in response.get_data(as_text=True)


def test_custom_settings_change_result(client):
    page = client.post("/analyze", data={
        "mode": "sequence", "sequence": webapp.EXAMPLE_SEQUENCE, "threshold": "2.1",
    }).get_data(as_text=True)
    assert "is present, but not in a hydrophobic region" in page


def test_export_record_workbook(client):
    response = client.get("/export/Q26473")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ["NCBI record", "Features", "Motif analysis", "KD profile", "TM prediction"]
    record = workbook["NCBI record"]
    fields = [row[0] for row in record.iter_rows(min_row=2, values_only=True)]
    assert fields[:4] == ["LOCUS", "DEFINITION", "ACCESSION", "VERSION"]
    assert "  AUTHORS" in fields and "ORIGIN" in fields
    origin_row = fields.index("ORIGIN") + 2
    assert record.cell(origin_row, 2).value.startswith("1 mraalvavaa")
    # no HTML should end up in the excel file
    assert not any("<" in str(c) for row in record.iter_rows(values_only=True) for c in row if c)


def test_export_analysis_for_pasted_sequence(client):
    response = client.post("/export-analysis", data={"sequence": webapp.EXAMPLE_SEQUENCE})
    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ["Motif analysis", "KD profile", "TM prediction"]
    assert workbook["KD profile"].max_row == 482


def test_health(client):
    assert client.get("/health").get_json() == {"status": "ok"}
