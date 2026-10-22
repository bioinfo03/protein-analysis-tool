"""Excel exports (openpyxl)."""

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ncbi
from analysis import kd_value

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True)
MOTIF_FILL = PatternFill("solid", fgColor="DCE9FB")
HIGH_FILL = PatternFill("solid", fgColor="FBE0DE")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
EXCEL_CELL_LIMIT = 32_000


def _cell_text(value):
    text = "" if value is None else str(value)
    if len(text) > EXCEL_CELL_LIMIT:
        return text[:EXCEL_CELL_LIMIT] + " ... [truncated: Excel cell limit]"
    return text


def _table(sheet, headers, rows, widths, wrap=True):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    if wrap:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = WRAP_TOP
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def add_record_sheets(workbook, record_text):
    """
    Sheet 1 has the GenPept record as Field / Details rows, in the same order
    as NCBI, with the ORIGIN sequence one line per row (like the example in
    the problem statement). Sheet 2 has the FEATURES table.
    """
    sheet = workbook.active
    sheet.title = "NCBI record"
    rows = []
    for entry in ncbi.parse_genpept(record_text):
        if entry.key == "FEATURES":
            rows.append(["FEATURES", "See the 'Features' sheet (one row per feature qualifier)."])
        elif entry.key == "ORIGIN":
            lines = ncbi.origin_lines(record_text)
            for index, line in enumerate(lines):
                rows.append(["ORIGIN" if index == 0 else "", line])
        else:
            rows.append(["  " * entry.level + entry.key, _cell_text(entry.value)])
    _table(sheet, ["Field", "Details"], rows, [16, 110])
    for row in sheet.iter_rows(min_row=2):
        if row[0].value and not row[0].value.startswith(" "):
            row[0].font = SECTION_FONT
        if (row[0].value == "ORIGIN") or (row[0].value == "" and row[1].value[:1].isdigit()):
            row[1].font = Font(name="Courier New")

    features = workbook.create_sheet("Features")
    feature_rows = []
    for number, feature in enumerate(ncbi.parse_features(record_text), start=1):
        qualifiers = feature.qualifiers or [("", "")]
        for name, value in qualifiers:
            feature_rows.append([number, feature.key, feature.location, name, _cell_text(value)])
    _table(features, ["#", "Feature", "Location", "Qualifier", "Value"], feature_rows, [6, 18, 22, 18, 90])


def add_analysis_sheets(workbook, analysis, source):
    s = analysis.settings
    summary = workbook.create_sheet("Motif analysis")
    rows = [
        ["Source", source],
        ["Sequence length", len(analysis.sequence)],
        ["Motif pattern", s.pattern],
        ["Residues allowed at '?'", s.allowed],
        ["Regular expression", analysis.pattern_regex],
        ["Window size", s.window],
        ["Threshold (mean KD score)", s.threshold],
        ["Motif occurrences", len(analysis.hits)],
        ["Occurrences in a high-hydrophobicity window", len(analysis.qualifying_hits)],
        ["Windows at or above threshold", analysis.high_window_count],
        ["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")],
    ]
    _table(summary, ["Parameter", "Value"], rows, [42, 60])

    summary.append([])
    header_row = summary.max_row + 1
    summary.append(["Motif", "Start", "End", "In high window?", "Best window", "Best window score",
                    "Windows passing threshold"])
    for cell in summary[header_row]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for hit in analysis.hits:
        best = hit.best_window
        summary.append([
            hit.motif, hit.start, hit.end, "Yes" if hit.qualifying else "No",
            f"{best.start}-{best.end}" if best else "-",
            best.score if best else None,
            ", ".join(f"{w.start}-{w.end} ({w.score:.2f})" for w in hit.high_windows) or "-",
        ])
        if hit.qualifying:
            for cell in summary[summary.max_row]:
                cell.fill = MOTIF_FILL
    if not analysis.hits:
        summary.append(["No occurrences of the motif were found."])
    for column, width in zip("ABCDEFG", [42, 60, 8, 16, 14, 18, 60]):
        summary.column_dimensions[column].width = width
    summary.auto_filter.ref = None

    profile = workbook.create_sheet("KD profile")
    window_by_center = {w.center: w for w in analysis.windows}
    motif_positions = {p for h in analysis.qualifying_hits for p in range(h.start, h.end + 1)}
    profile_rows = []
    for pos, aa in enumerate(analysis.sequence, start=1):
        w = window_by_center.get(pos)
        profile_rows.append([
            pos, aa, kd_value(aa),
            f"{w.start}-{w.end}" if w else "",
            w.score if w else None,
            ("Yes" if w.is_high else "No") if w else "",
            "Yes" if pos in motif_positions else "",
        ])
    _table(profile, ["Position", "Residue", "KD value", "Window centred here", "Window mean KD",
                     "At/above threshold", "In qualifying motif"], profile_rows,
           [10, 9, 10, 20, 16, 18, 18], wrap=False)
    for row in profile.iter_rows(min_row=2):
        if row[5].value == "Yes":
            row[4].fill = HIGH_FILL
        if row[6].value == "Yes":
            row[1].fill = MOTIF_FILL


def to_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def record_workbook(record_text, analysis=None, source=""):
    workbook = Workbook()
    add_record_sheets(workbook, record_text)
    if analysis is not None:
        add_analysis_sheets(workbook, analysis, source)
    return to_bytes(workbook)


def analysis_workbook(analysis, source):
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_analysis_sheets(workbook, analysis, source)
    return to_bytes(workbook)
